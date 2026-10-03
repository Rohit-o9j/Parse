"""Extraction layer: format detection, error handling, PDF and DOCX reading."""

from __future__ import annotations

import io
import zipfile

import pytest

from src.extraction import (
    CorruptedFileError,
    EmptyFileError,
    FileTooLargeError,
    NoExtractableTextError,
    UnsupportedFileTypeError,
    extract_docx,
    extract_pdf,
    extract_text,
    file_extension,
    validate_upload,
)
from src.settings import MAX_FILE_SIZE_BYTES


class TestValidation:
    def test_detects_pdf_from_signature(self, rahul_pdf_bytes: bytes) -> None:
        assert validate_upload(rahul_pdf_bytes, "whatever.pdf") == "pdf"

    def test_detects_docx_from_signature(self, rahul_docx_bytes: bytes) -> None:
        assert validate_upload(rahul_docx_bytes, "whatever.docx") == "docx"

    def test_signature_beats_a_wrong_extension(self, rahul_pdf_bytes: bytes) -> None:
        assert validate_upload(rahul_pdf_bytes, "resume.docx") == "pdf"

    def test_empty_file_rejected(self) -> None:
        with pytest.raises(EmptyFileError):
            validate_upload(b"", "resume.pdf")

    def test_oversized_file_rejected(self) -> None:
        with pytest.raises(FileTooLargeError):
            validate_upload(b"%PDF-" + b"0" * MAX_FILE_SIZE_BYTES, "resume.pdf")

    def test_unsupported_extension_rejected(self) -> None:
        with pytest.raises(UnsupportedFileTypeError):
            validate_upload(b"plain text content here", "resume.txt")

    def test_legacy_doc_rejected_with_guidance(self) -> None:
        with pytest.raises(UnsupportedFileTypeError) as info:
            validate_upload(b"\xd0\xcf\x11\xe0" + b"\x00" * 64, "resume.doc")
        assert ".docx" in info.value.hint

    def test_mislabelled_pdf_rejected_as_corrupt(self) -> None:
        with pytest.raises(CorruptedFileError):
            validate_upload(b"this is definitely not a pdf at all", "resume.pdf")

    def test_non_word_zip_rejected(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("xl/workbook.xml", "<workbook/>")
        with pytest.raises(CorruptedFileError):
            validate_upload(buffer.getvalue(), "sheet.docx")

    @pytest.mark.parametrize(
        ("name", "expected"),
        [("a.PDF", "pdf"), ("a.docx", "docx"), ("a", ""), ("", "")],
    )
    def test_file_extension(self, name: str, expected: str) -> None:
        assert file_extension(name) == expected


class TestPdfExtraction:
    def test_reads_expected_content(self, rahul_pdf_bytes: bytes) -> None:
        result = extract_pdf(rahul_pdf_bytes, "rahul_sharma.pdf")
        assert result.file_type == "pdf"
        assert result.page_count >= 1
        assert "RAHUL SHARMA" in result.text
        assert "rahul.sharma@example.com" in result.text
        assert "National Institute of Technology" in result.text
        assert result.metadata["extractor"] == "PyMuPDF"

    def test_corrupt_pdf_raises(self) -> None:
        with pytest.raises(CorruptedFileError):
            extract_pdf(b"%PDF-1.7\nbroken body with no xref", "broken.pdf")

    def test_textless_pdf_raises(self) -> None:
        import pymupdf

        document = pymupdf.open()
        document.new_page()
        payload = document.tobytes()
        document.close()

        with pytest.raises(NoExtractableTextError) as info:
            extract_pdf(payload, "blank.pdf")
        assert "scan" in info.value.hint.lower()


class TestDocxExtraction:
    def test_reads_expected_content(self, rahul_docx_bytes: bytes) -> None:
        result = extract_docx(rahul_docx_bytes, "rahul_sharma.docx")
        assert result.file_type == "docx"
        assert "RAHUL SHARMA" in result.text
        assert "Zeta Payments Pvt Ltd" in result.text
        assert result.metadata["paragraph_count"] > 10

    def test_corrupt_docx_raises(self) -> None:
        with pytest.raises(CorruptedFileError):
            extract_docx(b"PK\x03\x04 not really a docx", "broken.docx")

    def test_reads_tables_headers_and_hyperlinks(self) -> None:
        from docx import Document

        document = Document()
        document.sections[0].header.paragraphs[0].text = "Priya Nair - priya@example.com"
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skills"
        table.cell(0, 1).text = "Python, Terraform"
        table.cell(1, 0).text = "Location"
        table.cell(1, 1).text = "Kochi, Kerala, India"
        document.add_paragraph("Led the platform team for three years.")
        document.add_paragraph("Shipped weekly.", style="List Bullet")

        buffer = io.BytesIO()
        document.save(buffer)
        result = extract_docx(buffer.getvalue(), "table.docx")

        assert "Priya Nair" in result.text
        assert "Skills | Python, Terraform" in result.text
        assert "Kochi, Kerala, India" in result.text
        assert "- Shipped weekly." in result.text
        assert result.metadata["table_count"] == 1

    def test_empty_docx_raises(self) -> None:
        from docx import Document

        buffer = io.BytesIO()
        Document().save(buffer)
        with pytest.raises(NoExtractableTextError):
            extract_docx(buffer.getvalue(), "empty.docx")


class TestDispatcher:
    def test_dispatches_pdf_and_docx(
        self, rahul_pdf_bytes: bytes, rahul_docx_bytes: bytes
    ) -> None:
        assert extract_text(rahul_pdf_bytes, "a.pdf").file_type == "pdf"
        assert extract_text(rahul_docx_bytes, "a.docx").file_type == "docx"

    def test_warns_about_a_mislabelled_file(self, rahul_pdf_bytes: bytes) -> None:
        result = extract_text(rahul_pdf_bytes, "resume.docx")
        assert any("named '.docx'" in warning for warning in result.warnings)

    def test_pdf_and_docx_agree_on_content(
        self, rahul_pdf_bytes: bytes, rahul_docx_bytes: bytes
    ) -> None:
        """Both renderings of the same source must carry the same facts."""
        pdf_text = extract_text(rahul_pdf_bytes, "a.pdf").text
        docx_text = extract_text(rahul_docx_bytes, "a.docx").text
        for fragment in (
            "rahul.sharma@example.com",
            "+91 98765 43210",
            "Zeta Payments Pvt Ltd",
            "CGPA: 8.7/10",
            "Smart Resume Parser",
        ):
            assert fragment in pdf_text
            assert fragment in docx_text
